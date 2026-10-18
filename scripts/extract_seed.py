# Reads the Excel sheet and writes src/seed.json  (run: python scripts/extract_seed.py SEPTEMBER_26.xlsx)
import openpyxl, json, sys, re
f = sys.argv[1] if len(sys.argv) > 1 else 'SEPTEMBER_26.xlsx'
wb = openpyxl.load_workbook(f); wv = openpyxl.load_workbook(f, data_only=True)
YM = '2026-09'
def num(x): return float(x) if isinstance(x,(int,float)) else 0
varieties, bills, payments, employees = [], [], [], []
base = wb['RAJA']
names = {}
for r in range(3, 27):
    nm = re.sub(r'\s+\d*$', '', str(base.cell(r,3).value or '').strip()).strip()
    if not nm: continue
    vid = f'v{r-2}'
    varieties.append(dict(id=vid, name=nm, pcsPerBox=int(num(base.cell(r,4).value)) or 1,
        category='icecream', cost=num(base[f'AQ{r}'].value), empRate=num(base[f'AO{r}'].value), sellRate=num(base[f'AP{r}'].value)))
    names[r] = vid
# Idiyappam line (from the Jafer sheet): name, Rate (employee pays), Rate Work (customer), Rate Company (cost)
for k,(nm,er,sr,c) in enumerate([('IDIYAPPAM',.4,1,.4),('PUTTU',1,2,1),('KOLKATTAI',.4,1,.4),('THENGAI 3',1,0,0),('THENGAI 2.5',1,0,0)],1):
    varieties.append(dict(id=f'i{k}', name=nm, category='idiyappam', pcsPerBox=1, cost=c, empRate=er, sellRate=sr))
for sheet in ['RAJA', 'JAFER']:
    ws = wb[sheet]; eid = 'e_' + sheet.lower()
    employees.append(dict(id=eid, name=sheet.title()))
    for d in range(1, 32):
        col = 4 + d  # E = day 1
        items = []
        for r, vid in names.items():
            q = num(ws.cell(r, col).value)
            if q > 0:
                items.append(dict(varietyId=vid, category='icecream', qty=q, cost=num(ws[f'AQ{r}'].value), empRate=num(ws[f'AO{r}'].value), sellRate=num(ws[f'AP{r}'].value)))
        if items: bills.append(dict(id=f'b_{sheet.lower()}_{d}', date=f'{YM}-{d:02d}', employeeId=eid, items=items))
        pay = ws.cell(29, col).value
        if isinstance(pay,(int,float)) and pay > 0:
            payments.append(dict(id=f'p_{sheet.lower()}_{d}', date=f'{YM}-{d:02d}', employeeId=eid, amount=pay))
json.dump(dict(rates={}, varieties=varieties, employees=employees, bills=bills, payments=payments), open('src/seed.json','w'), indent=1)
print(len(varieties), 'varieties', len(bills), 'bills', len(payments), 'payments')
for v in varieties[:3]: print(v)
